"""PowerPoint renderer. Native tables/charts only. Layouts come from the branded master."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from . import style
from .data import REPO_ROOT
from .exhibits.base import Ctx, rgb


def _layout(prs, spec_tuple):
    name, idx = spec_tuple
    return next((l for l in prs.slide_layouts if l.name == name), prs.slide_layouts[idx])


def render(ctx: Ctx, out: Path) -> Path:
    spec, prov = ctx.spec, ctx.prov
    tmpl = REPO_ROOT / spec.template if spec.template else None
    branded = bool(tmpl and tmpl.exists())
    prs = Presentation(str(tmpl)) if branded else Presentation()
    if not branded:
        prs.slide_width, prs.slide_height = style.SLIDE_W, style.SLIDE_H

    t = prs.slides.add_slide(_layout(prs, style.MASTER_TITLE))
    t.shapes.title.text = spec.title
    sub = t.placeholders[1] if len(t.placeholders) > 1 else None
    if sub is not None:
        sub.text = f"{'  |  '.join(f'{k}: {v}' for k, v in prov.params.items())}  |  Owner: {spec.owner}"
    if not branded:
        t.shapes.title.left, t.shapes.title.top, t.shapes.title.width, t.shapes.title.height = Inches(1), Inches(2.4), Inches(11.3), Inches(1.5)
        if sub is not None:
            sub.left, sub.top, sub.width, sub.height = Inches(1), Inches(4.0), Inches(11.3), Inches(0.8)
    t.notes_slide.notes_text_frame.text = prov.notes()

    for s in spec.slides:
        slide = prs.slides.add_slide(_layout(prs, style.MASTER_CONTENT))
        slide.shapes.title.text = s.title
        if not branded:
            tt = slide.shapes.title
            tt.left, tt.top, tt.width, tt.height = Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9)
            tt.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
            tt.text_frame.paragraphs[0].runs[0].font.size = Pt(style.TITLE_PT)
        for ex, box in zip(s.exhibits, style.LAYOUTS[s.layout]):
            if s.layout == "two_up":
                cap = slide.shapes.add_textbox(box[0], box[1] - Inches(0.35), box[2], Inches(0.3))
                r = cap.text_frame.paragraphs[0].add_run()
                r.text, r.font.bold, r.font.size = ex.title, True, Pt(style.BODY_PT)
            ex.render_pptx(slide, box, ctx)
        foot = slide.shapes.add_textbox(*style.FOOTER)
        r = foot.text_frame.paragraphs[0].add_run()
        r.text, r.font.size, r.font.color.rgb = prov.footer(), Pt(style.FOOT_PT), rgb(style.MUTED)
        slide.notes_slide.notes_text_frame.text = prov.notes()

    path = out / f"{spec.id}_{prov.label}.pptx"
    prs.save(path)
    return path
