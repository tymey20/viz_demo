"""Reconcile a rendered deck to source. Runs on every build; non-zero exit on any problem."""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation

from .exhibits.base import Ctx

PLACEHOLDER = re.compile(r"lorem|ipsum|\bTODO\b|\[insert|\bxxx", re.I)


def validate_pptx(path: Path, ctx: Ctx) -> list[str]:
    spec = ctx.spec
    prs = Presentation(str(path))
    slides = list(prs.slides)
    problems = []
    if len(slides) != len(spec.slides) + 1:
        return [f"deck has {len(slides)} slides, spec expects {len(spec.slides) + 1}"]
    for n, (slide, s) in enumerate(zip(slides[1:], spec.slides), start=2):
        shapes = {
            "tables": [sh for sh in slide.shapes if sh.has_table],
            "charts": [sh for sh in slide.shapes if sh.has_chart],
            "text": [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame],
        }
        for ex in s.exhibits:
            problems += [f"slide {n}: {p}" for p in ex.reconcile(shapes, ctx)]
        if shapes["tables"] or shapes["charts"]:
            problems.append(f"slide {n}: unexpected extra exhibits in deck")
    for n, slide in enumerate(slides, start=1):
        if any(PLACEHOLDER.search(t) for t in (sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)):
            problems.append(f"slide {n}: placeholder text")
        if not slide.notes_slide.notes_text_frame.text.startswith("report_id="):
            problems.append(f"slide {n}: missing provenance in notes")
    return problems
