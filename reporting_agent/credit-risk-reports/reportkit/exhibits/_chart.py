"""Shared chart plumbing."""
from __future__ import annotations

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.util import Pt

from .. import style
from .base import rgb


def add_chart(slide, box, kind, categories, series: list[tuple[str, list[float]]], number_format: str):
    data = CategoryChartData()
    data.categories = categories
    for name, vals in series:
        data.add_series(name, vals)
    chart = slide.shapes.add_chart(kind, *box, data).chart
    chart.has_title = False
    chart.font.name, chart.font.size = style.FONT, Pt(style.LABEL_PT)
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    plot = chart.plots[0]
    plot.has_data_labels = len(series) <= 2
    if plot.has_data_labels:
        plot.data_labels.number_format, plot.data_labels.number_format_is_linked = number_format, False
        plot.data_labels.font.size = Pt(style.LABEL_PT)
    chart.value_axis.has_major_gridlines = False
    chart.value_axis.tick_labels.number_format, chart.value_axis.tick_labels.number_format_is_linked = number_format, False
    for i, s in enumerate(plot.series):
        s.format.fill.solid(); s.format.fill.fore_color.rgb = rgb(style.SERIES[i % len(style.SERIES)])
        try:
            s.format.line.color.rgb = rgb(style.SERIES[i % len(style.SERIES)])
        except Exception:
            pass
    return chart


def reconcile_chart(title, shapes, expected: list[list[float]], tol=1e-9):
    if not shapes["charts"]:
        return [f"{title}: no chart found on slide"]
    chart = shapes["charts"].pop(0).chart
    actual = [list(s.values) for s in chart.plots[0].series]
    if len(actual) != len(expected):
        return [f"{title}: {len(actual)} series in deck, {len(expected)} in source"]
    problems = []
    for si, (a, e) in enumerate(zip(actual, expected)):
        if len(a) != len(e):
            problems.append(f"{title}: series {si} has {len(a)} points, source {len(e)}")
            continue
        for pi, (x, y) in enumerate(zip(a, e)):
            if x is None or abs(x - y) > tol:
                problems.append(f"{title}: series {si} point {pi} shows {x}, source {y}")
    return problems
