from __future__ import annotations

import html
from typing import Literal

from pptx.util import Inches, Pt

from .. import style
from ..thresholds import rag
from .base import Ctx, Exhibit, rgb


class Table(Exhibit):
    type: Literal["table"]
    data: str
    key: str                      # first column
    columns: list[str]
    rag: list[str] = []           # columns colored by threshold

    def _rows(self, ctx: Ctx):
        df, spec = ctx.df(self.data), ctx.spec
        for _, row in df.iterrows():
            cells = [str(row[self.key])]
            for c in self.columns:
                cells.append(spec.fmt(c).text(row[c]))
            yield row, cells

    def render_pptx(self, slide, box, ctx):
        df, spec = ctx.df(self.data), ctx.spec
        cols = [self.key] + self.columns
        x, y, w, _ = box
        tbl = slide.shapes.add_table(len(df) + 1, len(cols), x, y, w, Inches(0.38) * (len(df) + 1)).table
        for j, c in enumerate(cols):
            tbl.cell(0, j).text = c.replace("_", " ").upper()
        for i, (row, cells) in enumerate(self._rows(ctx), start=1):
            for j, text in enumerate(cells):
                cell = tbl.cell(i, j)
                cell.text = text
                c = cols[j]
                if c in self.rag and c in spec.thresholds:
                    st = rag(float(row[c]), spec.thresholds[c], row)
                    cell.fill.solid(); cell.fill.fore_color.rgb = rgb(style.RAG[st])
                    cell.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(style.RAG_TEXT[st])
        for r in tbl.rows:
            for cell in r.cells:
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size, run.font.name = Pt(style.TABLE_PT), style.FONT

    def render_html(self, ctx):
        spec = ctx.spec
        cols = [self.key] + self.columns
        out = ["<tr>" + "".join(f"<th>{html.escape(c)}</th>" for c in cols) + "</tr>"]
        for row, cells in self._rows(ctx):
            tds = []
            for j, text in enumerate(cells):
                c, st = cols[j], ""
                if c in self.rag and c in spec.thresholds:
                    st = f' style="background:#{style.RAG[rag(float(row[c]), spec.thresholds[c], row)]}"'
                tds.append(f"<td{st}>{html.escape(text)}</td>")
            out.append("<tr>" + "".join(tds) + "</tr>")
        return "<table>" + "".join(out) + "</table>"

    def reconcile(self, shapes, ctx):
        problems = []
        if not shapes["tables"]:
            return [f"{self.title}: no table found on slide"]
        tbl = shapes["tables"].pop(0).table
        for i, (row, cells) in enumerate(self._rows(ctx), start=1):
            for j, expected in enumerate(cells):
                shown = tbl.cell(i, j).text
                if shown != expected:
                    problems.append(f"{self.title}: row {row[self.key]} col {([self.key]+self.columns)[j]} shows {shown!r}, source {expected!r}")
        return problems
