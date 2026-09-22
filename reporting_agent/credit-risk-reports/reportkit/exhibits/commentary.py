from __future__ import annotations

import html
from typing import Literal

from pptx.util import Pt

from .. import style
from .base import Exhibit, rgb


class CommentarySlide(Exhibit):
    type: Literal["commentary"]

    def render_pptx(self, slide, box, ctx):
        c = ctx.commentary
        tf = slide.shapes.add_textbox(*box).text_frame
        tf.word_wrap = True
        r = tf.paragraphs[0].add_run()
        r.text, r.font.bold, r.font.size = c.status, True, Pt(style.BODY_PT)
        r.font.color.rgb = rgb(style.RAG["G"] if c.signed_off_by else style.RAG["R"])
        for line in c.lines:
            p = tf.add_paragraph()
            run = p.add_run()
            run.text, run.font.size, run.font.name = line, Pt(style.BODY_PT), style.FONT
            p.space_before = Pt(6)

    def render_html(self, ctx):
        c = ctx.commentary
        return f"<p><b>{html.escape(c.status)}</b></p><ul>" + "".join(f"<li>{html.escape(l)}</li>" for l in c.lines) + "</ul>"

    def reconcile(self, shapes, ctx):
        txt = " ".join(shapes["text"])
        return [] if ctx.commentary.status in txt else [f"{self.title}: sign-off status missing from slide"]
