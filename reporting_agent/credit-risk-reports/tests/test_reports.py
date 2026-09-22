"""Every report must build and reconcile to source. New reports get added to REPORTS."""
import pytest

from reportkit import cli

REPORTS = [
    ("reports/pd_monitoring_quarterly", {"as_of": "2026-06-30"}),
    ("reports/raroc_committee", {"as_of": "2026-08-31"}),
]


def _args(cmd, report, params, out):
    return [cmd, report, "--out", str(out)] + [x for k, v in params.items() for x in ("-p", f"{k}={v}")]


@pytest.mark.parametrize("report,params", REPORTS)
def test_builds_and_reconciles(report, params, tmp_path):
    assert cli.main(_args("build", report, params, tmp_path)) == 0
    assert list(tmp_path.glob("*.pptx"))


def test_validator_catches_tampered_table(tmp_path):
    from pptx import Presentation
    from reportkit.validate import validate_pptx
    report, params = REPORTS[0]
    cli.main(_args("build", report, params, tmp_path))
    deck = next(tmp_path.glob("*.pptx"))
    prs = Presentation(str(deck))
    tbl = next(s for sl in prs.slides for s in sl.shapes if s.has_table).table
    tbl.cell(1, 2).text = "0.999"
    prs.save(str(deck))
    ctx = cli.prepare(report, params)
    assert any("shows '0.999'" in p for p in validate_pptx(deck, ctx))


def test_validator_catches_tampered_chart(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from reportkit.validate import validate_pptx
    report, params = REPORTS[1]
    cli.main(_args("build", report, params, tmp_path))
    deck = next(tmp_path.glob("*.pptx"))
    prs = Presentation(str(deck))
    chart = next(s for sl in prs.slides for s in sl.shapes if s.has_chart).chart
    cd = CategoryChartData(); cd.categories = ["a", "b"]; cd.add_series("x", [1.0, 2.0])
    chart.replace_data(cd)
    prs.save(str(deck))
    assert any("series" in p for p in validate_pptx(deck, cli.prepare(report, params)))
